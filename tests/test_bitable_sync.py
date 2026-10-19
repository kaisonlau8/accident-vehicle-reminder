"""飞书多维表格门店联系人同步测试（离线）。"""

import openpyxl

import bitable_sync
from bitable_sync import HEADERS, records_to_rows, write_list_xlsx


def _user(name: str, open_id: str) -> dict:
    return {"name": name, "id": open_id}


def _slot_fields(*, legacy: bool = False) -> dict:
    fields = {
        "序号": 1,
        "店代码": "D001",
        "店名称": "武汉测试店",
        "区域": "华中 1区",
        "督导": _user("督导甲", "ou_sup"),
        "督导手机号": "13800000000",
    }
    if legacy:
        fields["门店售后业务负责人"] = _user("旧负责人", "ou_old")
        fields["售后业务负责人手机号"] = "13800000009"
    else:
        fields.update({
            # 带空格与不带空格的字段名都要能识别。
            "门店售后业务负责人 1": _user("负责人甲", "ou_m1"),
            "售后业务负责人手机号 1": "138-0000-0001",
            "门店售后业务负责人2": _user("负责人乙", "ou_m2"),
            "售后业务负责人手机号2": 13800000002,
            "售后业务负责人手机号3": "",
            "门店售后业务负责人 4": _user("负责人甲", "ou_m1"),
            "售后业务负责人手机号 4": "13800000004",
        })
    return fields


def test_headers_append_manager_slots_after_district():
    assert HEADERS[10:13] == ["门店省", "门店市", "门店区"]
    assert HEADERS[13:] == [
        "售后负责人2", "售后电话2", "售后open_id2",
        "售后负责人3", "售后电话3", "售后open_id3",
        "售后负责人4", "售后电话4", "售后open_id4",
    ]
    assert len(HEADERS) == 22


def test_records_to_rows_parses_multi_slot_and_dedupes_by_open_id():
    rows, warnings = records_to_rows([_slot_fields()])

    assert warnings == []
    assert len(rows) == 1
    row = rows[0]
    assert len(row) == len(HEADERS)
    # 主列 = 第一个非空槽位。
    assert row[7:10] == ["负责人甲", "13800000001", "ou_m1"]
    # 2 号槽位。
    assert row[13:16] == ["负责人乙", "13800000002", "ou_m2"]
    # 3 号槽位为空。
    assert row[16:19] == ["", "", ""]
    # 4 号槽位与 1 号 open_id 相同，去重后不重复写入。
    assert row[19:22] == ["", "", ""]


def test_records_to_rows_supports_legacy_unnumbered_fields():
    rows, warnings = records_to_rows([_slot_fields(legacy=True)])

    assert warnings == []
    row = rows[0]
    assert row[7:10] == ["旧负责人", "13800000009", "ou_old"]
    assert row[13:22] == [""] * 9


def test_records_to_rows_primary_is_first_non_empty_slot():
    fields = {
        "店代码": "D002",
        "店名称": "测试店二",
        "门店售后业务负责人 2": _user("负责人乙", "ou_m2"),
        "售后业务负责人手机号2": "13800000002",
        "门店售后业务负责人 4": _user("负责人丁", "ou_m4"),
        "售后业务负责人手机号4": "13800000004",
    }
    rows, warnings = records_to_rows([fields])

    assert warnings == []
    assert rows[0][7:10] == ["负责人乙", "13800000002", "ou_m2"]
    assert rows[0][19:22] == ["负责人丁", "13800000004", "ou_m4"]


def test_records_to_rows_warns_when_all_slots_empty():
    rows, warnings = records_to_rows([{
        "店代码": "D003",
        "店名称": "无负责人门店",
        "门店售后业务负责人 1": [],
        "售后业务负责人手机号 1": "",
    }])

    assert len(rows) == 1
    assert rows[0][7:22] == [""] * 15
    assert warnings == ["D003 无负责人门店"]


def test_write_list_xlsx_writes_appended_headers(tmp_path, monkeypatch):
    path = tmp_path / "list.xlsx"
    monkeypatch.setattr(bitable_sync, "LIST_XLSX_PATH", path)
    rows, _ = records_to_rows([_slot_fields()])

    write_list_xlsx(rows, backup=False)

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    assert list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))) == HEADERS
    assert len(list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))) == len(HEADERS)


def test_slot1_numbered_and_legacy_not_mixed_and_dup_merge():
    import bitable_sync as bs
    fields = {
        "店代码": "S1",
        "门店售后业务负责人 1": [{"name": "新经理", "id": "ou_new"}],
        "门店售后业务负责人": [{"name": "旧经理", "id": "ou_old"}],
        "售后业务负责人手机号": "13900000000",
        "门店售后业务负责人 2": [{"name": "新经理", "id": "ou_new"}],
        "售后业务负责人手机号 2": "13800000000",
    }
    slots = bs._dedupe_manager_slots(bs._manager_slots(fields))
    people = [p for p in slots if p]
    assert people == [{"name": "新经理", "phone": "13800000000", "open_id": "ou_new"}]
